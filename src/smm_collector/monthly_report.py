"""月度业务报表 Excel 生成器（49 行主表 + 报价明细 + 计算说明）。

口径（与 portal_service 一致，由服务层计算后传入，本模块只负责版式）：
  - F 列月均价 = 依据采集报价计算的月均价（非 SMM 官方月均），
    仅来自数据库有效报价 average_price 的算术平均；
  - H 列环比 = (本月−上月)/上月，仅相邻两月均完整且上月均价非零时写入数值，
    否则留空并把原因写入 K 列备注（绝不写入假性 -100%）；
  - I 列预测为人工填写区，保留模板合并，不沿用旧预测、不自动编造；
  - 非 SMM 9 行保留原行与版式，价格/涨跌留空；
  - SMM 已合并规格（储能型≥2.50/2.40）无独立报价 → 月均价留空 + 备注说明；
  - 不覆盖用户原始模板与历史文件，每次生成独立工作簿。
"""
from __future__ import annotations

import io
import sqlite3
from datetime import date, timedelta
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from smm_collector import portal_service as ps

# 模板版式（与《市场价格数据模板》最新页一致）
COLUMN_WIDTHS = {"A": 10.3, "B": 9.4, "C": 10.2, "D": 9.6, "E": 36.6,
                 "F": 18.6, "G": 8.4, "H": 17.7, "I": 21.9, "J": 19.7, "K": 38.5}
PREDICTION_MERGES = ["I23:I25", "I41:I42", "I43:I44", "I45:I46"]
FONT_NAME = "微软雅黑"
HEADER_FILL = "1F3864"       # 深蓝表头
SECTION_FILL = "D9E2F3"      # 说明页小节底
FILL_LIGHT = "F2F6FB"        # 明细表头
BORDER_GRAY = "BFBFBF"

MONTHLY_SHEET_PREFIX = ""  # 主表名 = 月份（如 "8月"）


def _thin_border() -> Border:
    side = Side(style="thin", color=BORDER_GRAY)
    return Border(left=side, right=side, top=side, bottom=side)


def _header_style() -> tuple[Font, PatternFill, Alignment]:
    return (Font(name=FONT_NAME, size=12, bold=True, color="FFFFFF"),
            PatternFill(start_color=HEADER_FILL, end_color=HEADER_FILL, fill_type="solid"),
            Alignment(horizontal="center", vertical="center", wrap_text=True))


def _next_week_label(month: str) -> str:
    """下月第一个周一~周五 → 'MM.DD-MM.DD'（预测栏人工区标题）。"""
    bounds = ps.month_bounds(month)
    if bounds is None:
        return ""
    nxt = (bounds[1].replace(day=28) + timedelta(days=4)).replace(day=1)
    d = nxt
    while d.weekday() != 0:
        d += timedelta(days=1)
    return f"{d.strftime('%m.%d')}-{(d + timedelta(days=4)).strftime('%m.%d')}"


def _number_format_for(value: Decimal | None) -> str:
    if value is not None and abs(value) < 1:
        return "0.0000"
    return "#,##0.00"


def build_monthly_report(products: list[dict], con: sqlite3.Connection, month: str,
                         meta: dict | None = None) -> Workbook:
    """生成月度业务报表工作簿：主表 + 报价明细 + 计算说明。"""
    meta = meta or {}
    payload = ps.monthly_payload(products, con, month, meta)
    if "error" in payload:
        raise ValueError(payload["error"])
    rows, summary = payload["rows"], payload["summary"]
    by_id = {r["id"]: r for r in rows}
    first, last = ps.month_bounds(month)

    wb = Workbook()
    ws = wb.active
    ws.title = f"{int(month[5:7])}月"

    # ── 表头 ──
    headers = ["组织", "物料属性", "信息类别", "化学类型", "详细内容",
               f"月均价\n（{first.strftime('%m.%d')}-{last.strftime('%m.%d')}）",
               "单位", "涨跌\n（本月均价与上月均价相比）",
               f"下周价格预测\n（{_next_week_label(month)}）", "数据来源", "备注"]
    font, fill, align = _header_style()
    for ci, h in enumerate(headers, 1):
        c = ws.cell(row=1, column=ci, value=h)
        c.font, c.fill, c.alignment = font, fill, align
        c.border = _thin_border()
    ws.row_dimensions[1].height = 52

    # ── 数据行 ──
    data_font = Font(name=FONT_NAME, size=11)
    data_align_c = Alignment(horizontal="center", vertical="center", wrap_text=False)
    data_align_l = Alignment(horizontal="center", vertical="center", wrap_text=True)
    # 组织列：仅组首行写值，其余行由合并区域覆盖（与模板版式一致）
    ordered = sorted(rows, key=lambda x: x["template_row"])
    org_starts: set[int] = set()
    cur_org = None
    for r in ordered:
        if r["organization"] != cur_org:
            org_starts.add(r["template_row"])
            cur_org = r["organization"]
    for r in rows:
        row_idx = r["template_row"]
        rng = by_id[r["id"]]
        vals = {
            1: r["organization"] if row_idx in org_starts else None,
            2: r["material_attribute"], 3: r["info_category"],
            4: r["chemistry"], 5: r["template_detail"],
            6: float(Decimal(rng["monthly_avg"])) if rng["monthly_avg"] is not None else None,
            7: r["unit"],
            8: float(Decimal(rng["mom_pct"]) / 100) if rng["mom_pct"] is not None else None,
            9: None,  # 预测：人工填写区
            10: r["source"],
            11: _remark_for(rng),
        }
        for ci, v in vals.items():
            c = ws.cell(row=row_idx, column=ci, value=v)
            c.font = data_font
            c.border = _thin_border()
            c.alignment = data_align_l if ci == 5 else data_align_c
            if ci == 6 and v is not None:
                c.number_format = _number_format_for(Decimal(rng["monthly_avg"]))
            if ci == 8 and v is not None:
                # 涨红跌绿（Excel 自定义格式颜色码），与门户约定一致
                c.number_format = '[Red]+0.00%;[Green]-0.00%;0.00%'
        ws.row_dimensions[row_idx].height = 30 if len(r["template_detail"]) > 40 else 20

    # ── 组织列合并（按相邻业务行分组，与模板 A2:A22 / A23:A30 / A31:A40 / A41:A50 一致） ──
    start = end = None
    cur_org = None
    for r in ordered:
        if r["organization"] != cur_org:
            if cur_org is not None and end - start >= 1:
                ws.merge_cells(start_row=start, start_column=1, end_row=end, end_column=1)
            cur_org, start, end = r["organization"], r["template_row"], r["template_row"]
        else:
            end = r["template_row"]
    if cur_org is not None and end - start >= 1:
        ws.merge_cells(start_row=start, start_column=1, end_row=end, end_column=1)

    # ── 预测列合并（模板版式） ──
    for m in PREDICTION_MERGES:
        ws.merge_cells(m)

    # ── 列宽 / 冻结 ──
    for letter, w in COLUMN_WIDTHS.items():
        ws.column_dimensions[letter].width = w
    ws.freeze_panes = "A2"

    # ── 报价明细 ──
    ws_detail = wb.create_sheet("报价明细")
    detail_headers = ["业务行", "业务产品", "报价日期", "最低价", "最高价", "日均价",
                      "涨跌", "单位", "采集更新时间"]
    for ci, h in enumerate(detail_headers, 1):
        c = ws_detail.cell(row=1, column=ci, value=h)
        c.font = Font(name=FONT_NAME, size=10, bold=True, color="1F3864")
        c.fill = PatternFill(start_color=FILL_LIGHT, end_color=FILL_LIGHT, fill_type="solid")
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = _thin_border()
    dr = 2
    small = Font(name=FONT_NAME, size=10)
    for r in rows:
        pts_list = r["points"]
        if not pts_list:
            continue
        for p in pts_list:
            vals = [r["template_row"],
                    f"{r['display_name']}（{r['spec']}）" if r["spec"] else r["display_name"],
                    p["price_date"], p["min_price"], p["max_price"],
                    p["average_price"], p["change_value"], p["unit"], p["collected_at"]]
            for ci, v in enumerate(vals, 1):
                c = ws_detail.cell(row=dr, column=ci, value=v)
                c.font = small
                c.border = _thin_border()
                c.alignment = Alignment(horizontal="center", vertical="center")
            dr += 1
    for i, w in enumerate([8, 42, 12, 12, 12, 12, 10, 10, 21], 1):
        ws_detail.column_dimensions[get_column_letter(i)].width = w
    ws_detail.freeze_panes = "A2"

    # ── 计算说明 ──
    ws_note = wb.create_sheet("计算说明")
    ws_note.column_dimensions["A"].width = 30
    ws_note.column_dimensions["B"].width = 100
    note_font = Font(name=FONT_NAME, size=10)
    title_font = Font(name=FONT_NAME, size=11, bold=True, color="1F3864")
    r_idx = 1

    def section(text: str):
        nonlocal r_idx
        c = ws_note.cell(row=r_idx, column=1, value=text)
        c.font = title_font
        c.fill = PatternFill(start_color=SECTION_FILL, end_color=SECTION_FILL, fill_type="solid")
        r_idx += 1

    def kv(k: str, v: str):
        nonlocal r_idx
        ws_note.cell(row=r_idx, column=1, value=k).font = Font(name=FONT_NAME, size=10, bold=True)
        c = ws_note.cell(row=r_idx, column=2, value=v)
        c.font = note_font
        c.alignment = Alignment(wrap_text=True, vertical="top")
        r_idx += 1

    section("一、计算口径")
    kv("月均价定义", summary["method_note"])
    kv("月度区间", f"{summary['month_start']} ~ {summary['month_end']}"
                    + ("" if summary["month_ended"] else "（月份未结束，为截至当前的期间均价）"))
    kv("采集日历", f"日频采集日 {summary['calendar_daily_days']}/"
                   f"{summary['calendar_daily_weekdays']} 个工作日；"
                   f"周频报价 {summary['calendar_weekly_points']}/"
                   f"{summary['calendar_weekly_fridays']} 个周五")
    kv("完整性阈值", f"日频 ≥{meta.get('daily_completeness_ratio', 0.8):.0%}、"
                     f"周频 ≥{meta.get('weekly_completeness_ratio', 0.75):.0%}；"
                     "同一报价点重复采集不增权重（按报价日期去重）")
    kv("环比规则", "（本月月均价−上月月均价）÷上月月均价×100%；仅相邻两月均完整且上月月均价非零时计算")
    kv("日均价与月均价区别", "日均价=某一报价日期的 average_price 源字段；"
                             "月均价=当月有效报价 average_price 的算术平均（按有效发布日计，周频按当月发布报价点计）")
    kv("涨跌显示约定", "涨红跌绿，保留正负号")
    r_idx += 1

    section("二、逐行结果")
    for r in rows:
        mom = f"{r['mom_pct']}%" if r["mom_pct"] is not None else "—"
        kv(f"行{r['template_row']} {r['display_name']}",
           f"月均价={r['monthly_avg'] or '—'}（{r['n_points']}个报价点，"
           f"{r['date_from']}~{r['date_to']}）｜完整性={r['completeness']}：{r['reason']}"
           f"｜上月月均价={r['prev_monthly_avg'] or '—'}｜环比={mom}"
           + (f"｜环比未计算原因：{r['mom_reason']}" if r["mom_reason"] else "")
           + (f"｜口径说明：{r['note']}" if r["note"] else ""))
    r_idx += 1

    section("三、无独立报价 / 人工填写行")
    for r in rows:
        if r["mapping_status"] == "unavailable_merged":
            kv(f"行{r['template_row']} {r['display_name']}（{r['spec']}）", r["reason"])
        elif r["mapping_status"] == "manual":
            kv(f"行{r['template_row']} {r['display_name']}（来源：{r['source']}）",
               "非SMM来源，保留原行与版式，价格与涨跌留空供人工填写")
    r_idx += 1

    section("四、映射依据")
    kv("映射配置", "config/business_products.yaml（首页/走势/月报共用同一映射与数据服务层）")
    kv("LFP 价格点调整", "SMM 2026-07-17 正式生效：动力型≥2.55/2.50/2.40 与储能型≥2.30 "
                         "改名并改定义（新定义不含循环寿命承诺）；储能型≥2.50/2.40 并入对应密度价格点。"
                         "证据：docs/smm_lfp_price_point_change_2026.md")
    kv("映射核对", "40 条 SMM 业务行：34 条明确对应 + 4 条官方改名延续绑定 + 2 条无独立报价（留空）")

    return wb


def _remark_for(rng: dict) -> str | None:
    """主表 K 列备注：只写必要口径说明，干净行留空。"""
    parts = []
    if rng["mapping_status"] == "unavailable_merged":
        parts.append(rng["note"] or "暂无独立报价")
    if rng["mapping_status"] == "confirmed_renamed":
        parts.append("SMM 2026-07-17 新口径（不含循环寿命承诺）")
    if rng["mom_reason"]:
        parts.append(rng["mom_reason"])
    return "；".join(parts) or None


def build_monthly_report_xlsx(products: list[dict], con: sqlite3.Connection, month: str,
                              meta: dict | None = None) -> io.BytesIO:
    """生成并返回 xlsx 字节流（门户下载/脚本共用）。"""
    wb = build_monthly_report(products, con, month, meta)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def report_filename(month: str) -> str:
    return f"月度业务报表_{month}.xlsx"
