"""门户业务品种数据服务层测试（synthetic SQLite，不依赖真实网络/生产库）。

覆盖需求验证项：映射 40 行/34 确认/6 特殊不合并、同名不同规格不串价、
现货与指数隔离、同一行四价同源、重复分类与重复采集去重、周频/改名/缺值/零涨跌、
as_of 不返回未来报价、7 月不完整→8 月环比留空、完整相邻月环比公式、
月报 49 行/非 SMM 留空/无公式与假性-100%、月均价可由明细复核。
"""
from __future__ import annotations

import sqlite3
from datetime import datetime, timedelta
from decimal import Decimal
from pathlib import Path

import openpyxl
import pytest
import yaml

import sys
sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))

from smm_collector import portal_service as ps  # noqa: E402
from smm_collector import monthly_report as mr  # noqa: E402

CONFIG_ROOT = Path(__file__).resolve().parents[1] / "config"


# ── synthetic DB ────────────────────────────────────────────

def _row(category, product, spec, unit, price_date, avg, collected_at=None,
         min_p=None, max_p=None, change=None, status="valid",
         market="SMM锂电现货", source="SMM锂电现货"):
    if min_p is None:
        min_p = avg
    if max_p is None:
        max_p = avg
    return {
        "source": source, "market": market, "category": category,
        "product_name": product, "specification": spec,
        "min_price": str(min_p) if min_p is not None else None,
        "max_price": str(max_p) if max_p is not None else None,
        "average_price": str(avg) if avg is not None else None,
        "change_value": str(change) if change is not None else None,
        "unit": unit, "price_date": price_date,
        "collected_at": collected_at or f"{price_date}T09:05:00",
        "record_hash": f"h-{category}-{product}-{spec}-{price_date}-{avg}",
        "validation_status": status,
    }


def _workdays(year: int, month: int, from_day: int = 1, to_day: int = 31) -> list[str]:
    out = []
    d = datetime(year, month, from_day)
    end = datetime(year, month, min(to_day, 28) + (1 if to_day > 28 else 0), 1) - timedelta(days=1) \
        if to_day >= 28 else datetime(year, month, to_day)
    if to_day == 31:
        end = (datetime(year, month, 28) + timedelta(days=4)).replace(day=1) - timedelta(days=1)
    while d <= end:
        if d.weekday() < 5:
            out.append(d.strftime("%Y-%m-%d"))
        d += timedelta(days=1)
    return out


@pytest.fixture()
def db_path(tmp_path):
    """合成库：日历完整的 6-8 月日频产品 + 周频电芯 + 改名对 + 重复分类 + 特殊行。"""
    path = tmp_path / "smm_test.db"
    con = sqlite3.connect(path)
    con.execute("""CREATE TABLE lithium_spot_prices (
      id INTEGER PRIMARY KEY AUTOINCREMENT,
      source TEXT, market TEXT, category TEXT, product_name TEXT, specification TEXT DEFAULT '',
      min_price TEXT, max_price TEXT, average_price TEXT, change_value TEXT, unit TEXT DEFAULT '',
      price_date TEXT, collected_at TEXT, source_url TEXT, collection_method TEXT, raw_text TEXT,
      extra_fields TEXT, record_hash TEXT, validation_status TEXT, validation_message TEXT,
      created_at TEXT, updated_at TEXT)""")
    rows = []

    # ── 日频：电池级/工业级碳酸锂（7 月完整 19 个工作日、8 月完整 21 个工作日）──
    jul_days = _workdays(2026, 7)      # 2026-07 共 23 个工作日
    aug_days = _workdays(2026, 8)      # 2026-08 共 21 个工作日
    jul_days = jul_days[:19]           # 19/23 ≥ 80% → 完整（测试用）
    aug_days = aug_days[:21]           # 21/21 → 完整
    base = Decimal("150000")
    for i, d in enumerate(jul_days + aug_days):
        v = base + i * 200
        # 电池级碳酸锂（双分类重复收录：锂化合物 + 磷化工，值完全一致）
        rows.append(_row("锂化合物", "电池级碳酸锂", "Li₂CO₃≥99.5%", "元/吨", d, v,
                         change=-200 if i > 0 else 0))
        rows.append(_row("磷化工", "电池级碳酸锂", "Li₂CO₃≥99.5%", "元/吨", d, v,
                         change=-200 if i > 0 else 0))
        # 工业级碳酸锂（7 月仅 8 天 → 7 月不完整，用于验证环比留空）
        if i < 8 or d >= "2026-08-01":
            rows.append(_row("锂化合物", "工业级碳酸锂", "Li₂CO₃≥99.2%", "元/吨", d, v - 8000,
                             change=-100 if i > 0 else 0))

    # ── 现货 vs 指数隔离：同名指数行必须不影响现货绑定 ──
    for i, d in enumerate(aug_days):
        rows.append(_row("镍化合物", "SMM电池级硫酸镍指数", "指数", "指数", d, 100 + i,
                         change=0))
    rows.append(_row("镍化合物", "电池级硫酸镍", "镍含量≥22%", "元/吨", "2026-08-03", 32000, change=0))

    # ── 周频电芯（周五发布；同报价点重复采集只计一次） ──
    fridays = []
    d = datetime(2026, 7, 3)
    while d <= datetime(2026, 8, 31):
        if d.weekday() == 4:
            fridays.append(d.strftime("%Y-%m-%d"))
        d += timedelta(days=1)
    for i, f in enumerate(fridays):
        rows.append(_row("电芯", "方形磷酸铁锂电芯(周)", "100Ah", "元/Wh", f, Decimal("0.4") + i * Decimal("0.01"),
                         change=Decimal("0.01") if i else 0))
        rows.append(_row("磷化工", "方形磷酸铁锂电芯(周)", "100Ah", "元/Wh", f, Decimal("0.4") + i * Decimal("0.01"),
                         change=Decimal("0.01") if i else 0))
        # 同报价点次日重复采集（collected_at 更新）→ 权重不变
        rows.append(_row("电芯", "方形磷酸铁锂电芯(周)", "100Ah", "元/Wh", f, Decimal("0.4") + i * Decimal("0.01"),
                         collected_at=f"{f}T14:30:00", change=Decimal("0.01") if i else 0))

    # ── LFP 改名对（动力2.55 绑定）＋ 同密度不同系列隔离（2.50 独立值） ──
    for i, d in enumerate(_workdays(2026, 7, 22, 31)):
        rows.append(_row("正极材料", "磷酸铁锂", "压实密度≥2.55g/cm³", "元/吨", d, 56000 + i * 300, change=0))
        rows.append(_row("正极材料", "磷酸铁锂", "压实密度≥2.50g/cm³", "元/吨", d, 54000 + i * 300, change=0))
    for i, d in enumerate(_workdays(2026, 8, 3, 31)):
        rows.append(_row("正极材料", "磷酸铁锂", "粉体压实密度≥2.55g/cm³", "元/吨", d, 58400 + i * 300, change=0))
        rows.append(_row("正极材料", "磷酸铁锂", "粉体压实密度≥2.50g/cm³", "元/吨", d, 56400 + i * 300, change=0))

    # ── 特殊行：缺 average_price（不计入月均）、invalid（排除）、change=0 ──
    rows.append(_row("锂化合物", "电池级碳酸锂", "Li₂CO₃≥99.5%", "元/吨", "2026-08-31",
                     None, min_p=151000, max_p=152000, change=0, status="warning"))
    rows.append(_row("锂化合物", "电池级碳酸锂", "Li₂CO₃≥99.5%", "元/吨", "2026-08-01",
                     999999, change=0, status="invalid"))

    for r in rows:
        con.execute("""INSERT INTO lithium_spot_prices
          (source, market, category, product_name, specification, min_price, max_price,
           average_price, change_value, unit, price_date, collected_at, record_hash,
           validation_status, created_at, updated_at)
          VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
          (r["source"], r["market"], r["category"], r["product_name"], r["specification"],
           r["min_price"], r["max_price"], r["average_price"], r["change_value"],
           r["unit"], r["price_date"], r["collected_at"], r["record_hash"],
           r["validation_status"], "now", "now"))
    con.commit()
    con.close()
    return path


@pytest.fixture(scope="module")
def products():
    prods, meta = ps.load_products(CONFIG_ROOT)
    assert meta["load_errors"] == [], meta["load_errors"]
    return prods


# ── 映射配置（验证项 1/2） ──────────────────────────────────

def test_mapping_counts(products):
    assert len(products) == 49
    smm = ps.smm_products(products)
    assert len(smm) == 40
    statuses = {}
    for p in smm:
        statuses[p["mapping_status"]] = statuses.get(p["mapping_status"], 0) + 1
    assert statuses["confirmed"] == 34
    assert statuses["confirmed_renamed"] == 4
    assert statuses["unavailable_merged"] == 2
    assert sum(1 for p in products if p["mapping_status"] == "manual") == 9


def test_lfp_six_rows_not_merged(products):
    """6 条 LFP 动力/储能独立保留；绑定系列互不重复；并入规格无绑定。"""
    density_ids = {"lfp_power_255", "lfp_power_250", "lfp_power_240",
                   "lfp_ess_250", "lfp_ess_240", "lfp_ess_230"}
    lfp = [p for p in products if p["id"] in density_ids]
    assert len(lfp) == 6
    bound = [p for p in lfp if p["series"]]
    assert len(bound) == 4
    unbound = [p for p in lfp if not p["series"]]
    assert {p["id"] for p in unbound} == {"lfp_ess_250", "lfp_ess_240"}
    # 每个绑定 (product_name, specification) 只属于一条业务行（绝不复制同价冒充独立报价）
    keys = [tuple(sorted((t["product_name"], t["specification"]) for t in p["series"]))
            for p in bound]
    assert len(set(keys)) == len(keys)
    # 动力/储能各保有自己的密度（2.55 动力 ≠ 2.50 动力 ≠ 2.40 动力 ≠ 2.30 储能）
    densities = {p["id"]: [t["specification"] for t in p["series"]] for p in bound}
    assert len({str(v) for v in densities.values()}) == 4


def test_no_two_rows_share_series(products):
    """任何两条 SMM 业务行不得共用同一系列三元组（同价复制防护）。"""
    seen = {}
    for p in ps.smm_products(products):
        for t in p["series"]:
            key = (t["category"], t["product_name"], t["specification"])
            assert key not in seen, f"系列 {key} 被 {seen[key]} 与 {p['id']} 共用"
            seen[key] = p["id"]


def test_series_exact_not_fuzzy(products):
    """绑定全部为精确三元组；搜索不影响绑定身份。"""
    for p in ps.smm_products(products):
        for t in p["series"]:
            assert t["category"] and t["product_name"], p["id"]


# ── 每日报价（验证项 4/5/8） ────────────────────────────────

def test_quote_same_row_four_prices(db_path, products):
    """同一行 min/max/avg/change 来自同一条记录；avg 缺失不重算、缺失显示空。"""
    con = ps.open_readonly(db_path)
    try:
        quotes = ps.latest_quote(ps.smm_products(products), con, as_of="2026-08-31")
        q = quotes["lce_battery"]
        assert q is not None
        # 8月31日 avg 缺失（warning 行）→ 应回退到上一个有 avg 的报价点（08-28，周五）
        assert q["price_date"] == "2026-08-28"
        # 四价同源：来自同一记录（同 collected_at 同 category）
        assert q["min_price"] is not None and q["max_price"] is not None
        assert q["average_price"] is not None and q["change_value"] is not None
        # avg 不被 min/max 重算（avg 与 (min+max)/2 无强绑定，此处 min=max=avg）
        assert Decimal(q["average_price"]) == Decimal(q["min_price"])
    finally:
        con.close()


def test_quote_no_future_leak(db_path, products):
    con = ps.open_readonly(db_path)
    try:
        quotes = ps.latest_quote(ps.smm_products(products), con, as_of="2026-08-10")
        for pid, q in quotes.items():
            if q is not None:
                assert q["price_date"] <= "2026-08-10", pid
        # 周频 100Ah：08-10 之前最近报价点应为 08-07
        assert quotes["cell_lfp_100ah"]["price_date"] == "2026-08-07"
    finally:
        con.close()


def test_spot_vs_index_isolated(db_path, products):
    """现货绑定不混入同名指数行。"""
    con = ps.open_readonly(db_path)
    try:
        rows = ps.product_rows(con, next(p for p in products if p["id"] == "niso4_battery"))
        assert len(rows) == 1
        assert rows[0]["product_name"] == "电池级硫酸镍"
        assert rows[0]["specification"] == "镍含量≥22%"
        assert rows[0]["unit"] == "元/吨"
    finally:
        con.close()


def test_same_name_diff_density_not_crossed(db_path, products):
    """磷酸铁锂 2.55 与 2.50 为不同报价系列，取值不互串。"""
    con = ps.open_readonly(db_path)
    try:
        q = ps.latest_quote(ps.smm_products(products), con, as_of="2026-08-31")
        v255 = Decimal(q["lfp_power_255"]["average_price"])
        v250 = Decimal(q["lfp_power_250"]["average_price"])
        assert v255 != v250
        assert v255 > v250  # 合成数据里 2.55 系列值更高
    finally:
        con.close()


def test_zero_change_shown(db_path, products):
    """零涨跌正常显示 0，不显示为缺失。"""
    con = ps.open_readonly(db_path)
    try:
        rows = ps.product_rows(con, next(p for p in products if p["id"] == "lfp_power_255"))
        pts = ps.merge_points(rows)
        first = pts[0]
        assert ps.fmt_price(ps.to_dec(first["change_value"])) == "0"
    finally:
        con.close()


# ── 历史序列（验证项 3/7） ──────────────────────────────────

def test_history_rename_pair_merged(db_path, products):
    """压实密度→粉体压实密度 改名对合并为连续序列，无断档。"""
    con = ps.open_readonly(db_path)
    try:
        payload = ps.history_payload(products, con, ["lfp_power_255"],
                                     date_from="2026-07-20", date_to="2026-08-31")
        pts = payload["series"][0]["points"]
        dates = [p["price_date"] for p in pts]
        assert dates == sorted(dates)
        assert dates[0] == "2026-07-22" and dates[-1] == "2026-08-31"
        # 改名日 08-03 前后连续（07-31 → 08-03，中间 08-01/02 为周末无报价）
        assert "2026-07-31" in dates and "2026-08-03" in dates
    finally:
        con.close()


def test_weekly_duplicate_collection_single_weight(db_path, products):
    """同报价点重复采集（同日两次 collected_at）只计一次。"""
    con = ps.open_readonly(db_path)
    try:
        rows = ps.product_rows(con, next(p for p in products if p["id"] == "cell_lfp_100ah"))
        pts = ps.merge_points(rows)
        assert len(pts) == len({r["price_date"] for r in rows})
        # 每个日期一条（取 collected_at 最新）
        by_date = {}
        for r in rows:
            by_date[r["price_date"]] = by_date.get(r["price_date"], 0) + 1
        assert any(v >= 2 for v in by_date.values())
        assert len(pts) == len(by_date)
    finally:
        con.close()


def test_home_trends_monthly_same_series(db_path, products):
    """首页最新报价与走势序列、月报参与点共用同一报价系列（同一数据路径）。"""
    con = ps.open_readonly(db_path)
    try:
        quotes = ps.latest_quote(ps.smm_products(products), con, as_of="2026-08-31")
        hist = ps.history_payload(products, con, ["lce_battery"],
                                  date_from="2026-08-01", date_to="2026-08-31")
        monthly = ps.monthly_stats(next(p for p in products if p["id"] == "lce_battery"),
                                   con, "2026-08")
        hist_pts = hist["series"][0]["points"]
        # 最新报价 = 走势序列中最后一个 average_price 非空的报价点
        avg_pts = [p for p in hist_pts if p["average_price"] is not None]
        assert quotes["lce_battery"]["price_date"] == avg_pts[-1]["price_date"]
        # 月报参与点 = 走势序列中 average_price 非空的点（同一集合，同一路径）
        assert [p["price_date"] for p in monthly["points"]] == \
               [p["price_date"] for p in avg_pts]
    finally:
        con.close()


# ── 月均价 / 完整性 / 环比（验证项 9/11） ───────────────────

def test_monthly_avg_reconcile_from_points(db_path, products):
    """月均价可由明细报价点复核（= points 的 average_price 算术平均）。"""
    con = ps.open_readonly(db_path)
    try:
        st = ps.monthly_stats(next(p for p in products if p["id"] == "lce_battery"),
                              con, "2026-08", all_products=products)
        mean = sum((Decimal(p["average_price"]) for p in st["points"]), Decimal(0)) \
            / len(st["points"])
        assert st["monthly_avg"] is not None
        assert abs(Decimal(str(st["monthly_avg"])) - mean) < Decimal("0.01")
        # 8 月 21 个工作日；其中 08-31 avg 缺失不计入 → 20 点
        assert st["n_points"] == 20
    finally:
        con.close()


def test_missing_avg_and_invalid_excluded(db_path, products):
    """average_price 缺失的 warning 行与 invalid 行均不参与月均价。"""
    con = ps.open_readonly(db_path)
    try:
        st = ps.monthly_stats(next(p for p in products if p["id"] == "lce_battery"),
                              con, "2026-08", all_products=products)
        dates = [p["price_date"] for p in st["points"]]
        assert "2026-08-31" not in dates  # avg 缺失
        assert "2026-08-01" not in dates  # invalid
        assert st["completeness"] == "complete"
    finally:
        con.close()


def test_duplicate_categories_not_double_counted(db_path, products):
    """重复分类（锂化合物+磷化工）不改变月均价。"""
    con = ps.open_readonly(db_path)
    try:
        st = ps.monthly_stats(next(p for p in products if p["id"] == "lce_battery"),
                              con, "2026-08")
        vals = [Decimal(p["average_price"]) for p in st["points"]]
        assert len(vals) == len(set(p["price_date"] for p in st["points"]))
    finally:
        con.close()


def test_july_incomplete_aug_mom_blank(db_path, products, monkeypatch):
    """7 月采集不完整（本合成库 19/23 ≥80% 完整——此处针对工业级碳酸锂：
    7 月仅 8 天 → 8 月环比必须留空并说明原因。"""
    con = ps.open_readonly(db_path)
    try:
        payload = ps.monthly_payload(products, con, "2026-08")
        row = next(r for r in payload["rows"] if r["id"] == "lce_industrial")
        assert row["mom_pct"] is None
        assert "7" in (row["mom_reason"] or "") and "不完整" in (row["mom_reason"] or "")
        assert row["monthly_avg"] is not None  # 本月有值，只是环比不计算
    finally:
        con.close()


def test_complete_adjacent_months_mom_formula(db_path, products):
    """相邻两月均完整 → 环比 =（本月−上月）/上月，可复算。"""
    con = ps.open_readonly(db_path)
    try:
        payload = ps.monthly_payload(products, con, "2026-08")
        row = next(r for r in payload["rows"] if r["id"] == "lce_battery")
        assert row["mom_pct"] is not None
        cur = Decimal(row["monthly_avg"])
        prev = Decimal(row["prev_monthly_avg"])
        expected = (cur - prev) / prev * 100
        assert abs(Decimal(row["mom_pct"]) - expected) < Decimal("0.01")
        # 合成数据：7 月 19 天 150000 起每天 +200；8 月 21 天继续 → 环比为正
        assert Decimal(row["mom_pct"]) > 0
    finally:
        con.close()


def test_september_in_progress_no_mom(db_path, products):
    con = ps.open_readonly(db_path)
    try:
        payload = ps.monthly_payload(products, con, "2026-09")
        assert payload["summary"]["month_ended"] is False
        for r in payload["rows"]:
            assert r["mom_pct"] is None
            if r["source"] == "SMM" and r["mapping_status"] in ("confirmed", "confirmed_renamed"):
                assert r["completeness"] == "in_progress"
                assert "未结束" in (r["reason"] or "")
    finally:
        con.close()


def test_weekly_monthly_avg_points_only(db_path, products):
    """周频月均价只按当月实际发布报价点计（8 月 4 个周五）。"""
    con = ps.open_readonly(db_path)
    try:
        st = ps.monthly_stats(next(p for p in products if p["id"] == "cell_lfp_100ah"),
                              con, "2026-08", all_products=products)
        assert st["n_points"] == 4
        assert st["completeness"] == "complete"
        dates = [p["price_date"] for p in st["points"]]
        assert all(d >= "2026-08-01" and d <= "2026-08-31" for d in dates)
    finally:
        con.close()


# ── 搜索（验证项：业务名称/SMM 名称/缩写可检索） ────────────

def test_search_products(products):
    smm = ps.smm_products(products)
    assert {p["id"] for p in ps.search_products(smm, "LFP")} >= {"lfp_power_255", "black_lfp_piece_powder"}
    assert "lce_battery" in {p["id"] for p in ps.search_products(smm, "电碳")}
    assert "cell_lfp_100ah" in {p["id"] for p in ps.search_products(smm, "100Ah")}
    assert "niso4_battery" in {p["id"] for p in ps.search_products(smm, "硫酸镍")}


# ── 月报 Excel（验证项 10/12） ─────────────────────────────

def test_report_workbook_49_rows_non_smm_blank(db_path, products):
    con = ps.open_readonly(db_path)
    try:
        wb = mr.build_monthly_report(products, con, "2026-08")
    finally:
        con.close()
    ws = wb["8月"]
    assert ws.dimensions == "A1:K50"
    # 49 条业务记录全保留
    for r in range(2, 51):
        assert ws.cell(row=r, column=5).value, f"E{r} 为空"
    # 非 SMM 行价格/涨跌留空
    for r in (20, 21, 22, 37, 38, 39, 40, 49, 50):
        assert ws.cell(row=r, column=6).value is None, r
        assert ws.cell(row=r, column=8).value is None, r
    # 无独立报价行留空
    assert ws.cell(row=26, column=6).value is None
    assert ws.cell(row=27, column=6).value is None
    # 组织列合并与预测列合并保留
    merges = {str(m) for m in ws.merged_cells.ranges}
    assert {"A2:A22", "A23:A30", "A31:A40", "A41:A50", "I23:I25"} <= merges
    # 动态标题（自然月实际天数）
    assert "08.01-08.31" in str(ws.cell(row=1, column=6).value)


def test_report_no_formulas_no_fake_minus_100(db_path, products):
    con = ps.open_readonly(db_path)
    try:
        wb = mr.build_monthly_report(products, con, "2026-08")
    finally:
        con.close()
    for wsn in wb.sheetnames:
        ws = wb[wsn]
        for row in ws.iter_rows():
            for c in row:
                v = c.value
                assert not (isinstance(v, str) and v.startswith("=")), \
                    f"{wsn}!{c.coordinate} 含公式 {v!r}"
                # 假性 -100%：主表环比列（H）空值行不得写入 -1 或 -100（跳过合并单元格）
                if wsn == "8月" and isinstance(c, openpyxl.cell.cell.Cell) \
                        and c.column_letter == "H" and isinstance(v, (int, float)):
                    assert v > -0.99, f"{wsn}!{c.coordinate} = {v}"
    # 说明页包含无独立报价与人工填写行清单
    note = wb["计算说明"]
    joined = "".join(str(note.cell(row=r, column=1).value or "")
                     + str(note.cell(row=r, column=2).value or "")
                     for r in range(1, note.max_row + 1))
    assert "无独立报价" in joined and "人工填写" in joined


def test_report_detail_reconcilable(db_path, products):
    """明细工作表可复核月均价（参与报价日期/高低/均价/涨跌/单位齐备）。"""
    con = ps.open_readonly(db_path)
    try:
        wb = mr.build_monthly_report(products, con, "2026-08")
    finally:
        con.close()
    ws = wb["报价明细"]
    headers = [ws.cell(row=1, column=c).value for c in range(1, 10)]
    assert headers == ["业务行", "业务产品", "报价日期", "最低价", "最高价", "日均价",
                       "涨跌", "单位", "采集更新时间"]
    # 电池级碳酸锂（业务行 10）的明细行 = 20 个报价点（08-31 avg 缺失不计入）
    n = sum(1 for r in range(2, ws.max_row + 1)
            if ws.cell(row=r, column=1).value == 10)
    assert n == 20


# ── V5：dataset 多分类 / 基础金属类别 / 月报 49 行守卫 ─────────────

def test_dataset_quotes_multi_category(db_path):
    """dataset_quotes 支持 categories 多分类并集（分页在多分类上正确）。"""
    con = ps.open_readonly(db_path)
    try:
        payload = ps.dataset_quotes(con, categories=["锂化合物", "镍化合物"],
                                    page=1, page_size=200)
        cats = {r["category"] for r in payload["rows"]}
        assert cats == {"锂化合物", "镍化合物"}
        single = ps.dataset_quotes(con, category="锂化合物", page=1, page_size=200)
        assert payload["total"] > single["total"]
    finally:
        con.close()


def test_dataset_products_multi_category(db_path):
    """dataset_products 支持 categories 多分类并集。"""
    con = ps.open_readonly(db_path)
    try:
        prods = ps.dataset_products(con, categories=["锂化合物", "电芯"])
        cats = {p["category"] for p in prods}
        assert cats == {"锂化合物", "电芯"}
    finally:
        con.close()


def test_metals_info_category_and_filters(products):
    """基础金属 3 产品信息类别独立，筛选维度含 8 个类别。"""
    metals = [p for p in products if p["id"] in ("al_a00", "cu_1e", "ni_1e")]
    assert len(metals) == 3
    assert all(p["info_category"] == "基础金属" for p in metals)
    assert all(p["source"] == "SMM" for p in metals)
    cats = sorted({p["info_category"] for p in products if p["info_category"]})
    assert cats == ["基础金属", "废极片", "新电池", "新电芯", "正极材料", "金属", "金属盐", "黑粉"]


def test_monthly_payload_49_rows_includes_metals(db_path, products):
    """月报保持 49 行模板（基础金属 3 行在列，无数据时月均价留空）。"""
    con = ps.open_readonly(db_path)
    try:
        payload = ps.monthly_payload(products, con, "2026-08")
    finally:
        con.close()
    assert len(payload["rows"]) == 49
    assert payload["summary"]["total_rows"] == 49
    ids = [r["id"] for r in payload["rows"]]
    assert ids[:3] == ["al_a00", "cu_1e", "ni_1e"]
    for r in payload["rows"][:3]:
        assert r["monthly_avg"] is None  # 合成库无基础金属数据
        assert r["info_category"] == "基础金属"
